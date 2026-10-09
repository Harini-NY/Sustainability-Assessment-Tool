import os
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


def set_impact_colour(cell,impact):
    impact = str(impact).strip()
    if impact == "+":
        colour = RGBColor(0,176,80) # Green
    elif impact == "=":
        colour = RGBColor(255,192,0) # Yellow
    elif impact == "-":
        colour = RGBColor(237, 125, 49) # Orange
    elif impact == "- -":
        colour = RGBColor(255, 0, 0)  # Red
    elif impact.upper() == "N/A":
        colour = RGBColor(191, 191, 191)  # Grey
    else:
        colour = RGBColor(255,255,255)  # White
    cell.fill.solid()
    cell.fill.fore_color.rgb = colour

def format_cell(cell, text, font_size=10, bold=False):  # Formatting Function
    cell.text = str(text)

    text_frame = cell.text_frame
    text_frame.word_wrap = True
    text_frame.auto_size = None
    

    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER

        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
            run.font.bold = bold


def generate_ppt(      # receives results from streamlit # opens template.pptx
    parameter_results,
    systems_results,
    tool_completed,
    design_reviewed,
    key_takeaway,
    site,
    project_title
):

    prs = Presentation("Templates/template.pptx")

    slide = prs.slides[0]

    
    # Find all tables
    tables = []

    for i, shape in enumerate(slide.shapes):
        if shape.has_table:
            tables.append(shape.table)
   
            

    parameter_table = tables[0]
    systems_table = tables[2]
    assessment_table = tables[1]
    design_table = tables[3]

    for row in systems_table.rows:
        row.height = Pt(20)

    format_cell(assessment_table.cell(0,1), tool_completed)
    format_cell(design_table.cell(0,1), design_reviewed) 


    # PARAMETERS # Fills the table # Cells change colour accordingly 
 

    parameter_start_row = 1

    for i, item in enumerate(parameter_results):

        row = parameter_start_row + i

        if row >= len(parameter_table.rows):
            break

        format_cell(
            parameter_table.cell(row, 0),
            f"{item['Parameter']},{item['Unit']}"
        )
        format_cell(parameter_table.cell(row, 1), item["Baseline"])
        format_cell(parameter_table.cell(row, 2), item["Projected Load"])
        format_cell(parameter_table.cell(row, 3), item["Tariff"])

        format_cell(parameter_table.cell(row, 4), item["Impact"])
        set_impact_colour(parameter_table.cell(row,4), item["Impact"])

        format_cell(parameter_table.cell(row, 5), item["Mitigation"])

        format_cell(parameter_table.cell(row, 6), item["New Impact"])
        set_impact_colour(parameter_table.cell(row, 6), item["New Impact"])


    # SYSTEMS
    

    systems_start_row = 1
    print("SYSTEM TABLE ROWS =", len(systems_table.rows))
    tables = []

    for i, item in enumerate(systems_results):

        row = systems_start_row + i
        if row >= len(systems_table.rows):
            break
        format_cell(
            systems_table.cell(row, 0),
            f"{item["System"]},{item["Unit"]}"
        )
        # Site Baseline
        format_cell(systems_table.cell(row, 1),item["Baseline"])
        # Projected
        format_cell(systems_table.cell(row, 2),item["Projected Load"])
        # Redundancy Limit
        format_cell(systems_table.cell(row, 3),item["Redundancy"])
        # Redundancy Impact
        format_cell(systems_table.cell(row, 4),item["Impact"])
        set_impact_colour(systems_table.cell(row, 4),item["Impact"])
        # Mitigation
        format_cell(systems_table.cell(row, 5),item["Mitigation"])
        # New Impact
        format_cell(systems_table.cell(row, 6),item["New Impact"])
        set_impact_colour(systems_table.cell(row, 6),item["New Impact"])



    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue

        existing_text = shape.text.strip()
        if "Key Takeaway" in existing_text:
            shape.text = f"Key Takeaway\n{key_takeaway}"
        elif "Utilities and Sustainability Assessment" in existing_text:
            shape.text = (f"Utilities and Sustainability Assessment")

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(12)

    output_file = f"saved_assessments/{project_title}.pptx"
    prs.save(output_file)
    return output_file